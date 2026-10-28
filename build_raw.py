#!/usr/bin/env python3
"""
build_raw.py — Generate the deliberately messy synthetic premium bordereau
and the machine-readable error ledger.

Scenario (100% synthetic, invented names throughout):
  Coverholder : Harbourline Insurance Services Ltd (invented), PIN 123456782
  Binder      : UMR B0876HLB250001X, Agreement BND-2024-1187
                Incepted 1 Jul 2024, 18-month term to 31 Dec 2025
                => Year of Account (CR0010) = 2024 for EVERY policy under it
                   (Lloyd's definition: YoA = year the binding authority incepted)
  Sections    : 1 Commercial Property   risk codes {B5}      commission 25.0%
                2 Household             risk codes {HP, B4}   commission 20.0%
                3 Employers' Liability  risk codes {W3}      commission 17.5%
  Period      : 1-31 December 2025 (premium bordereau)
  Territory   : GB risks in GBP, IE risks in EUR; all settled in GBP.
                FX (guide convention: rate = original units per 1 settlement
                unit; settlement amount = original / rate): EUR->GBP 1.16
  Taxes       : GB risks IPT 12% of gross; IE risks nil in this sample.
                Net Premium to London = Gross - Commission - Taxes.

Idempotent: fixed seed, overwrites its outputs.
Outputs: raw-bordereau.csv, raw-bordereau.xlsx, error-ledger.csv
"""
import csv, json, random, os
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = os.path.dirname(os.path.abspath(__file__))
random.seed(20251231)

# ---------------------------------------------------------------- contract ---
UMR = "B0876HLB250001X"
AGREEMENT_NO = "BND-2024-1187"
COVERHOLDER = "Harbourline Insurance Services Ltd"
PIN = "123456782"
BROKER_REF = "HLB"                      # invented sample value
YOA = 2024                              # binder incepted 1 Jul 2024
PERIOD_START = date(2025, 12, 1)
PERIOD_END = date(2025, 12, 31)
FX_EUR_GBP = 1.16                       # 1 GBP = 1.16 EUR (guide convention)

SECTIONS = {
    1: {"class": "Commercial Property",   "codes_gb": ["B5"], "codes_ie": ["B5"], "comm": 25.0},
    2: {"class": "Household",             "codes_gb": ["HP"], "codes_ie": ["B4"], "comm": 20.0},
    3: {"class": "Employers' Liability",  "codes_gb": ["W3"], "codes_ie": ["W3"], "comm": 17.5},
}

COMPANIES = [
    "Ashford Retail Group Ltd", "Beaconfield Manufacturing Ltd", "Cartwright Logistics Ltd",
    "Dunmore Hotels Ltd", "Ellerby Construction Ltd", "Fenwick Foods Ltd", "Grangewood Joinery Ltd",
    "Hollis Print & Packaging Ltd", "Ingleby Farms Ltd", "Jessop Engineering Ltd",
    "Kestrel Components Ltd", "Larkfield Care Homes Ltd", "Malloway Textiles Ltd",
    "Northgate Vehicle Hire Ltd", "Otterbourne Leisure Ltd", "Penkridge Plastics Ltd",
    "Quarry Bank Estates Ltd", "Ravensworth Roofing Ltd", "Stonebridge Kitchens Ltd",
    "Thornbury Motors Ltd", "Ullswater Marinas Ltd", "Ventnor Fisheries Ltd",
    "Westcombe Bakery Group Ltd", "Yardley Storage Ltd", "Zetland Office Services Ltd",
    "Aldergrove Inns Ltd", "Bramhope Dental Group Ltd", "Charnwood Garden Centres Ltd",
    "Dovedale Clothing Ltd", "Eastmere Electrical Ltd",
]
IE_COMPANIES = [
    "Liffey Retail Partners Ltd", "Connacht Food Services Ltd", "Boyne Logistics DAC",
    "Galway Bay Hotels Ltd", "Shannon Components Ltd", "Cork Harbour Storage Ltd",
]
PEOPLE = [
    ("James", "Hartley"), ("Sarah", "M Ellison"), ("David", "Okafor"), ("Emma", "R Whitfield"),
    ("Michael", "Donnelly"), ("Lucy", "Pemberton"), ("Thomas", "R Ashworth"), ("Priya", "Nair"),
    ("George", "Holloway"), ("Anna", "Kowalska"), ("Peter", "S Lindqvist"), ("Mary", "Gallagher"),
    ("John", "P Braithwaite"), ("Kate", "Middleton-Rees"), ("Andrew", "Frost"), ("Helen", "M Vardy"),
    ("Robert", "Chalmers"), ("Julia", "Nowak"), ("Mark", "T Ellery"), ("Susan", "Pritchard"),
    ("Kevin", "Doyle"), ("Laura", "B Ashcroft"), ("Stephen", "Moss"), ("Rachel", "T Pemberton"),
    ("Brian", "Kavanagh"), ("Claire", "Donoghue"), ("Paul", "Renshaw"), ("Fiona", "McAllister"),
]
IE_PEOPLE = [("Sean", "Murphy"), ("Aoife", "Byrne"), ("Liam", "O Connor"), ("Niamh", "Kelly"),
             ("Ciaran", "Walsh"), ("Siobhan", "Reilly")]
SECTORS = ["Retail", "Manufacturing", "Hospitality", "Construction", "Logistics & Distribution",
           "Professional Services", "Agriculture", "Healthcare", "Education", "Technology",
           "Food & Beverage", "Property Management"]
GB_COUNTIES = ["Kent", "Surrey", "Lancashire", "West Yorkshire", "Devon", "Norfolk", "Cheshire",
               "Staffordshire", "Hampshire", "Warwickshire", "Northumberland", "Somerset"]
IE_COUNTIES = ["Dublin", "Cork", "Galway", "Limerick", "Waterford"]

DATE_FORMATS = ["%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d.%m.%Y"]

def fmt_date(d, style):
    return d.strftime(DATE_FORMATS[style % len(DATE_FORMATS)])

def money_str(v, ccy, style):
    s = f"{v:,.2f}"
    if style == 0 and ccy == "GBP":
        return "£" + s
    if style == 0 and ccy == "EUR":
        return "€" + s
    return s

def pct_str(p, style):
    if style == 0:
        return f"{p:g}%"
    if style == 1:
        return f"{p:g}"
    return f"{p/100:.4f}".rstrip("0").rstrip(".")

def r2(v):
    return round(v + 1e-9, 2)

# ------------------------------------------------------------ base records ---
# Build 165 original-premium records + 35 endorsement records = 200 rows.
records = []
cert_seq = 0

def next_cert():
    global cert_seq
    cert_seq += 1
    return f"HLB/2024/{cert_seq:04d}"

def make_original(i):
    # section mix: ~45% S1, ~35% S2, ~20% S3
    r = random.random()
    section = 1 if r < 0.45 else (2 if r < 0.80 else 3)
    ie = random.random() < 0.15 and section in (1, 2)
    country = "IE" if ie else "GB"
    ccy = "EUR" if ie else "GBP"
    code = random.choice(SECTIONS[section]["codes_ie" if ie else "codes_gb"])
    # inception: weighted to Dec 2025, some earlier in binder term
    if random.random() < 0.62:
        inc = date(2025, 12, random.randint(1, 28))
    else:
        start = date(2024, 7, 1)
        inc = start + timedelta(days=random.randint(0, 515))
        if inc > date(2025, 11, 30):
            inc = date(2025, 11, random.randint(1, 30))
    exp = inc + timedelta(days=364)
    if section == 2:
        if ie:
            fn, ln = random.choice(IE_PEOPLE)
            sub = random.choice(IE_COUNTIES)
        else:
            fn, ln = random.choice(PEOPLE)
            sub = random.choice(GB_COUNTIES)
        insured, first, sector = f"{fn} {ln}", fn, ""
    else:
        insured = random.choice(IE_COMPANIES if ie else COMPANIES)
        first = ""
        sector = random.choice(SECTORS)
        sub = random.choice(IE_COUNTIES if ie else GB_COUNTIES)
    if section == 1:
        gross = r2(random.uniform(800, 45000))
        ssum = r2(gross * random.uniform(40, 120))
    elif section == 2:
        gross = r2(random.uniform(150, 2500))
        ssum = r2(random.uniform(60000, 450000))
    else:
        gross = r2(random.uniform(900, 28000))
        ssum = 0
    comm_pct = SECTIONS[section]["comm"]
    tax = r2(gross * 0.12) if country == "GB" else 0.0
    return {
        "section": section, "class": SECTIONS[section]["class"], "risk_code": code,
        "country": country, "ccy": ccy, "insured": insured, "first": first,
        "sector": sector, "subdiv": sub, "inception": inc, "expiry": exp,
        "sum_insured": ssum, "gross": gross, "comm_pct": comm_pct, "tax": tax,
        "trans_type": "Original Premium", "trans_eff": inc,
        "new_ren": random.choice(["New", "Renewal"]), "cert": next_cert(),
    }

originals = [make_original(i) for i in range(165)]
records.extend(originals)

# endorsements reference existing certificates
for j in range(35):
    base = random.choice(originals)
    rec = dict(base)
    frac = random.uniform(0.05, 0.40)
    if random.random() < 0.65:
        rec["trans_type"] = "Additional Premium"
        rec["gross"] = r2(base["gross"] * frac)
    else:
        rec["trans_type"] = "Return Premium"
        rec["gross"] = -r2(base["gross"] * frac)
    rec["tax"] = r2(rec["gross"] * 0.12) if base["country"] == "GB" else 0.0
    rec["trans_eff"] = date(2025, 12, random.randint(1, 30))
    rec["new_ren"] = ""
    records.append(rec)

random.shuffle(records)
for idx, rec in enumerate(records, start=1):
    rec["row"] = idx
    comm_amt = r2(rec["gross"] * rec["comm_pct"] / 100.0)
    net = r2(rec["gross"] - comm_amt - rec["tax"])
    rec["comm_amt"] = comm_amt
    rec["net"] = net
    if rec["ccy"] == "GBP":
        rec["fx"] = ""
        rec["net_settle"] = net
    else:
        rec["fx"] = FX_EUR_GBP
        rec["net_settle"] = r2(net / FX_EUR_GBP)

# ------------------------------------------------------------- error plant ---
ledger = []
eid = 0

def add_ledger(rec, field, eclass, label, raw_val, expected, disposition, basis):
    global eid
    eid += 1
    ledger.append({
        "error_id": f"E{eid:03d}", "raw_row": rec["row"], "cert_ref": rec["cert_raw"],
        "field": field, "error_class": eclass, "class_label": label,
        "raw_value": raw_val, "expected_correct_value": expected,
        "disposition": disposition, "rule_basis": basis,
    })

used = set()
def pick(pred, n, pool=None):
    cands = [r for r in (pool or records) if r["row"] not in used and pred(r)]
    random.shuffle(cands)
    chosen = cands[:n]
    for r in chosen:
        used.add(r["row"])
    return chosen

# raw (as-keyed) working copies of the fields that get planted on
for rec in records:
    rec["umr_raw"] = UMR
    rec["yoa_raw"] = str(YOA)
    rec["section_raw"] = str(rec["section"])
    rec["code_raw"] = rec["risk_code"]
    rec["cert_raw"] = rec["cert"]
    rec["ccy_raw"] = rec["ccy"]
    rec["sector_raw"] = rec["sector"]
    rec["inc_raw"] = rec["inception"]
    rec["exp_raw"] = rec["expiry"]
    rec["comm_amt_raw"] = rec["comm_amt"]
    rec["net_raw"] = rec["net"]
    rec["net_settle_raw"] = rec["net_settle"]

# C1 — wrong year of account (policy year stamped instead of binder YoA)
for rec in pick(lambda r: True, 10):
    bad = "2025" if rec["row"] % 2 == 0 else "2023"
    rec["yoa_raw"] = bad
    add_ledger(rec, "Year of Account", "C1", "Wrong year of account", bad, str(YOA),
               "CORRECTED",
               "CR0010: YoA = year the binding authority incepted (2024), not the policy year")

# C2 — risk code missing (derivable from section + location of risk)
for rec in pick(lambda r: True, 8):
    rec["code_raw"] = ""
    add_ledger(rec, "Risk Code", "C2", "Risk code missing", "", rec["risk_code"],
               "CORRECTED",
               "CR0016: code derived from the binder section's authorised code set + location of risk")

# C3 — risk code malformed (format) — 7 correctable + 1 not-a-code (flagged)
mangle = {"B5": ["b5", "B-5", "B 5"], "HP": ["hp", "HP "], "B4": ["B_4"], "W3": ["w3", "W-3"]}
mangled_pool = pick(lambda r: True, 7)
for rec in mangled_pool:
    bad = random.choice(mangle[rec["risk_code"]])
    rec["code_raw"] = bad
    add_ledger(rec, "Risk Code", "C3", "Risk code malformed", bad, rec["risk_code"],
               "CORRECTED",
               "Lloyd's Risk Code Guidance: codes are alphanumeric; format normalised and validated against the section's authorised set")
for rec in pick(lambda r: True, 1):
    rec["code_raw"] = "ZZ"
    add_ledger(rec, "Risk Code", "C3", "Risk code not a Lloyd's code", "ZZ", "",
               "FLAGGED_FOR_CLIENT",
               "Lloyd's Risk Code Guidance: 'ZZ' is not a Lloyd's risk code and cannot be derived — coverholder to confirm correct code")

# C4 — UMR malformed (hand-keyed; must equal contract UMR and meet format rule)
umr_bad = ["b0876hlb250001x", "B0876-HLB250001X", "B0876 HLB250001X", "B0876HLB250010X",
           "B0876HLB25001X", "B0876HLB250001XX", "B0877HLB250001X", "B0876HLB250001"]
for rec, bad in zip(pick(lambda r: True, 8), umr_bad):
    rec["umr_raw"] = bad
    add_ledger(rec, "Unique Market Reference (UMR)", "C4", "UMR malformed / not matching contract",
               bad, UMR, "CORRECTED",
               "CR0005 + user guide UMR format (starts 'B', alphanumeric only, max 17 chars); single-binder file — corrected to the contract UMR")

# C5 — section blank on a multi-section binder
for rec in pick(lambda r: True, 7):
    rec["section_raw"] = ""
    add_ledger(rec, "Section No", "C5", "Section number blank", "", str(rec["section"]),
               "CORRECTED",
               "CR0007 + v5.2 (Bulletin Y5261): section reference required on multi-section binders; derived from class of business")

# C6 — duplicate certificate reference on Original Premium rows (flagged)
orig_pool = [r for r in records if r["trans_type"] == "Original Premium"]
for rec in pick(lambda r: True, 5, pool=orig_pool):
    donor = random.choice([r for r in orig_pool if r["row"] != rec["row"]])
    rec["cert_raw"] = donor["cert"]
    add_ledger(rec, "Certificate Ref", "C6", "Duplicate certificate reference",
               donor["cert"], "", "FLAGGED_FOR_CLIENT",
               "CR0029: certificate reference must be unique per risk; the correct reference cannot be invented — coverholder to confirm")

# C7 — premium arithmetic: (a) commission amount, (b) net, (c) settlement net
for rec in pick(lambda r: True, 4):
    bad = r2(rec["comm_amt"] + random.choice([137.50, -92.25, 250.00, -48.75]))
    delta_net = r2(rec["comm_amt"] - bad)          # keyed net followed the wrong commission
    rec["comm_amt_raw"] = bad
    rec["net_raw"] = r2(rec["net"] + delta_net)
    if rec["ccy"] != "GBP":
        rec["net_settle_raw"] = r2(rec["net_raw"] / FX_EUR_GBP)
    else:
        rec["net_settle_raw"] = rec["net_raw"]
    add_ledger(rec, "Commission Amount", "C7", "Commission amount does not equal gross x commission %",
               f"{bad:.2f}", f"{rec['comm_amt']:.2f}", "CORRECTED",
               "CR0061/CR0062: commission amount recomputed as gross premium x commission %")
for rec in pick(lambda r: True, 3):
    bad = r2(rec["net"] + random.choice([500.00, -275.40, 118.90]))
    rec["net_raw"] = bad
    rec["net_settle_raw"] = r2(bad / FX_EUR_GBP) if rec["ccy"] != "GBP" else bad
    add_ledger(rec, "Net Premium to London (Original Currency)", "C7",
               "Net premium does not reconcile (gross - commission - taxes)",
               f"{bad:.2f}", f"{rec['net']:.2f}", "CORRECTED",
               "CR0065: net recomputed as gross premium paid this time - commission amount - total taxes and levies")
fx_pool = [r for r in records if r["ccy"] != "GBP"]
for rec in pick(lambda r: True, 2, pool=fx_pool):
    bad = r2(rec["net_settle"] + random.choice([64.20, -31.55]))
    rec["net_settle_raw"] = bad
    add_ledger(rec, "Net Premium to London (Settlement Currency)", "C7",
               "Settlement-currency net does not equal original net / rate of exchange",
               f"{bad:.2f}", f"{rec['net_settle']:.2f}", "CORRECTED",
               "CR0067/CR0068: settlement net recomputed as original-currency net / rate of exchange")

# C8 — industry sector blank on commercial risks (mandatory for commercial, v5.2)
comm_pool = [r for r in records if r["section"] in (1, 3)]
for rec in pick(lambda r: True, 6, pool=comm_pool):
    rec["sector_raw"] = ""
    add_ledger(rec, "Insured Trade or Industry Sector", "C8",
               "Industry sector blank on a commercial risk", "", "",
               "FLAGGED_FOR_CLIENT",
               "v5.2 / Bulletin Y5261: trade or industry of the insured is mandatory for commercial risks; cannot be invented — coverholder to supply")

# C9 — currency code invalid: 6 format variants (correctable) + 1 indeterminate (flagged)
ccy_bad = {"GBP": ["GBp", "£", "gbp"], "EUR": ["EURO", "Eur", "€"]}
for rec in pick(lambda r: True, 6):
    bad = random.choice(ccy_bad[rec["ccy"]])
    rec["ccy_raw"] = bad
    add_ledger(rec, "Original Currency", "C9", "Currency code not ISO 4217 format",
               bad, rec["ccy"], "CORRECTED",
               "CR0020 + user guide: ISO 4217 three-letter codes; unambiguous symbol/abbreviation normalised")
for rec in pick(lambda r: True, 1):
    rec["ccy_raw"] = "—"
    add_ledger(rec, "Original Currency", "C9", "Currency indeterminate in source file",
               "—", "", "FLAGGED_FOR_CLIENT",
               "CR0020: original currency cannot be determined from the source file — coverholder to confirm")

# C10 — date logic: expiry on or before inception (flagged; policy dates are source facts)
for rec in pick(lambda r: True, 6):
    bad_exp = rec["inception"] - timedelta(days=random.choice([1, 45, 200, 364]))
    rec["exp_raw"] = bad_exp
    add_ledger(rec, "Risk Expiry Date", "C10", "Expiry date on or before inception date",
               bad_exp.isoformat(), "", "FLAGGED_FOR_CLIENT",
               "CR0030/CR0031: expiry must follow inception; the correct period of cover is a source-record fact — coverholder to confirm (not silently altered)")

# ---------------------------------------------------------------- raw file ---
HEADERS = ["Coverholder", "UMR / Binder Ref", "Agreement No.", "Period From", "Period To",
           "Sect", "Class", "Risk Cd", "Ins Type", "YoA", "Cert Ref", "Insured Name",
           "Insured First Name", "Trade / Sector", "Insured Country", "Insured County/State",
           "New/Ren", "Inception", "Expiry", "Sum Insured", "Risk Country",
           "Risk County/State", "Trans Type", "Trans Eff Date", "Curr", "Gross Prem (this time)",
           "Annual Gross", "Terr Prem", "Comm %", "Comm Amt", "Tax Type", "Tax Jurisdiction",
           "Tax Amt", "Net to London", "Settle Curr", "FX Rate", "Net to London (Settle Ccy)"]
COUNTRY_RAW = {"GB": ["United Kingdom", "UK", "Great Britain", "GB"],
               "IE": ["Ireland", "IE", "Eire"]}

def raw_row(rec, i):
    style = i % 4
    country_txt = random.choice(COUNTRY_RAW[rec["country"]])
    annual = rec["gross"] if rec["trans_type"] == "Original Premium" else ""
    terr = r2(rec["gross"] * 0.02) if (rec["country"] == "GB" and rec["section"] == 1
                                       and rec["trans_type"] == "Original Premium") else ""
    return [
        COVERHOLDER, rec["umr_raw"], AGREEMENT_NO, PERIOD_START.isoformat(), PERIOD_END.isoformat(),
        rec["section_raw"], rec["class"], rec["code_raw"], "Direct", rec["yoa_raw"],
        rec["cert_raw"], rec["insured"], rec["first"], rec["sector_raw"], country_txt,
        rec["subdiv"], rec["new_ren"], fmt_date(rec["inc_raw"], style), fmt_date(rec["exp_raw"], style + 1),
        (f"{rec['sum_insured']:,.2f}" if rec["sum_insured"] else ""),
        country_txt, rec["subdiv"], rec["trans_type"],
        (fmt_date(rec["trans_eff"], style) if rec["trans_type"] != "Original Premium" else ""),
        rec["ccy_raw"], money_str(rec["gross"], rec["ccy"], style % 2),
        (f"{annual:,.2f}" if annual != "" else ""), (f"{terr:,.2f}" if terr != "" else ""),
        pct_str(rec["comm_pct"], style % 3), f"{rec['comm_amt_raw']:,.2f}",
        ("IPT" if rec["country"] == "GB" else ""), ("United Kingdom" if rec["country"] == "GB" else ""),
        f"{rec['tax']:,.2f}", f"{rec['net_raw']:,.2f}", "GBP",
        (f"{rec['fx']:.2f}" if rec["fx"] != "" else ""), f"{rec['net_settle_raw']:,.2f}",
    ]

title_rows = [
    ["HARBOURLINE INSURANCE SERVICES LTD — PREMIUM BORDEREAU (as keyed by coverholder)"],
    [f"Binder UMR: {UMR}   Agreement: {AGREEMENT_NO}   Period: December 2025   Year of Account: {YOA}"],
    [],
]
rows = [raw_row(r, i) for i, r in enumerate(records)]
tot_gross = r2(sum(r["gross"] for r in records))
tot_comm = r2(sum(r["comm_amt_raw"] for r in records))
tot_tax = r2(sum(r["tax"] for r in records))
tot_net = r2(sum(r["net_raw"] for r in records))
tot_net_s = r2(sum(r["net_settle_raw"] for r in records))
totals_row = ["TOTALS"] + [""] * 24 + [f"{tot_gross:,.2f}", "", "", "", f"{tot_comm:,.2f}",
                                       "", "", f"{tot_tax:,.2f}", f"{tot_net:,.2f}", "", "",
                                       f"{tot_net_s:,.2f}"]

with open(os.path.join(OUT, "raw-bordereau.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    for t in title_rows:
        w.writerow(t)
    w.writerow(HEADERS)
    w.writerows(rows)
    w.writerow(totals_row)

wb = Workbook()
ws = wb.active
ws.title = "Premium Bordereau (raw)"
for t in title_rows:
    ws.append(t)
ws.append(HEADERS)
for c in ws[4]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="D9D9D9")
for r in rows:
    ws.append(r)
ws.append(totals_row)
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
wb.save(os.path.join(OUT, "raw-bordereau.xlsx"))

with open(os.path.join(OUT, "error-ledger.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(ledger[0].keys()))
    w.writeheader()
    w.writerows(sorted(ledger, key=lambda e: (e["raw_row"], e["error_id"])))

n_corr = sum(1 for e in ledger if e["disposition"] == "CORRECTED")
n_flag = sum(1 for e in ledger if e["disposition"] == "FLAGGED_FOR_CLIENT")
print(f"rows={len(records)} planted_errors={len(ledger)} corrected={n_corr} flagged={n_flag}")
print(f"keyed totals: gross={tot_gross:,.2f} comm={tot_comm:,.2f} tax={tot_tax:,.2f} "
      f"net={tot_net:,.2f} net_settle={tot_net_s:,.2f}")
